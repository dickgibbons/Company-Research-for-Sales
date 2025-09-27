#!/usr/bin/env python3
"""
EPAM Business Intelligence Web UI
Provides a web interface to analyze companies and generate downloadable Word reports
"""

import os
import sys
import json
import tempfile
from datetime import datetime
from typing import Dict
import importlib.util

# Flask imports
from flask import Flask, render_template, request, jsonify, send_file, flash, redirect, url_for
import io
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn

# Import the master agent chain
def import_master_chain():
    """Dynamic import of master agent chain"""
    spec = importlib.util.spec_from_file_location(
        "master_chain",
        "/Users/dickgibbons/master_agent_chain.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.MasterEPAMIntelligenceChain

MasterEPAMIntelligenceChain = import_master_chain()

app = Flask(__name__)
app.secret_key = 'epam-intelligence-key-2025'

@app.route('/')
def index():
    """Main page with company input form"""
    return render_template('index.html')

@app.route('/analyze', methods=['POST'])
def analyze_company():
    """Analyze company and return results"""
    try:
        ticker = request.form.get('ticker', '').strip().upper()

        if not ticker:
            flash('Please enter a valid ticker symbol', 'error')
            return redirect(url_for('index'))

        # Get optional API keys
        grok_key = request.form.get('grok_key', '').strip() or None
        perplexity_key = request.form.get('perplexity_key', '').strip() or None

        # Initialize and run analysis
        master_chain = MasterEPAMIntelligenceChain(grok_key, perplexity_key)
        result = master_chain.analyze_company_complete(ticker)

        if result['status'] == 'success':
            # Store result in session or temporary file for download
            temp_file = tempfile.NamedTemporaryFile(mode='w', suffix='.json', delete=False)
            json.dump(result, temp_file, indent=2, default=str)
            temp_file.close()

            return render_template('results.html',
                                 result=result,
                                 ticker=ticker,
                                 temp_file=temp_file.name)
        else:
            flash(f'Analysis failed: {result.get("error", "Unknown error")}', 'error')
            return redirect(url_for('index'))

    except Exception as e:
        flash(f'System error: {str(e)}', 'error')
        return redirect(url_for('index'))

@app.route('/download/<path:temp_file>')
def download_report(temp_file):
    # Handle double slash issue
    if temp_file.startswith('/'):
        temp_file = temp_file[1:]
    """Generate and download Word document report"""
    try:
        # Load the analysis result
        with open(temp_file, 'r') as f:
            result = json.load(f)

        # Generate Word document
        doc = create_word_report(result)

        # Save to temporary file
        temp_docx = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
        doc.save(temp_docx.name)
        temp_docx.close()

        # Clean up JSON temp file
        try:
            os.unlink(temp_file)
        except:
            pass

        # Generate filename
        ticker = result.get('ticker', 'UNKNOWN')
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'EPAM_Analysis_{ticker}_{timestamp}.docx'

        return send_file(
            temp_docx.name,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

    except Exception as e:
        flash(f'Error generating report: {str(e)}', 'error')
        return redirect(url_for('index'))

def create_word_report(result: Dict) -> Document:
    """Create a comprehensive Word document report"""
    doc = Document()

    # Title page
    title = doc.add_heading('EPAM BUSINESS INTELLIGENCE REPORT', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Company info
    doc.add_heading(f'Target Company: {result.get("company_name", "Unknown")} ({result.get("ticker", "N/A")})', 1)
    doc.add_paragraph(f'Analysis Date: {result.get("timestamp", "Unknown")[:10]}')
    doc.add_paragraph(f'Intelligence Sources: SEC Filings + Strategic Analysis + Social/News Monitoring')

    # Executive Summary
    doc.add_heading('EXECUTIVE SUMMARY', 1)

    # SEC Filing Intelligence
    agent1_result = result.get('agent_1_result', {})
    filings = agent1_result.get('filings', {})
    doc.add_paragraph(f'• SEC Filings Analyzed: {len(filings)} documents')
    for form_type, filing in filings.items():
        doc.add_paragraph(f'  - {form_type}: Filed {filing.get("filing_date", "Unknown")}', style='List Bullet')

    # Strategic Opportunities
    agent2_result = result.get('agent_2_result', {})
    total_opps = agent2_result.get('total_opportunities', 0)
    high_priority = agent2_result.get('high_priority_count', 0)
    doc.add_paragraph(f'• Strategic Opportunities: {total_opps} identified ({high_priority} high priority)')

    # Social Intelligence
    agent3_result = result.get('agent_3_result', {})
    twitter_count = len(agent3_result.get('twitter_intelligence', []))
    exec_count = len(agent3_result.get('executive_intelligence', []))
    doc.add_paragraph(f'• Market Intelligence: {twitter_count} social mentions, {exec_count} executive sources')

    # Strategic Opportunities Detail
    if agent2_result.get('opportunities'):
        doc.add_heading('TOP EPAM OPPORTUNITIES', 1)
        for i, opp in enumerate(agent2_result['opportunities'][:5], 1):
            doc.add_heading(f'{i}. {opp["opportunity_area"]} [{opp["priority_level"]} Priority]', 2)
            doc.add_paragraph(f'EPAM Services: {opp["epam_services"]}')
            doc.add_paragraph(f'Confidence Score: {opp["confidence_score"]}/1.0')
            doc.add_paragraph(f'Evidence Count: {opp["evidence_count"]} indicators')
            doc.add_paragraph('Key Indicators:')
            for indicator in opp.get('key_indicators', [])[:3]:
                doc.add_paragraph(f'• {indicator}', style='List Bullet')
            doc.add_paragraph('')

    # Social Intelligence Detail
    if agent3_result.get('twitter_intelligence'):
        doc.add_heading('SOCIAL MEDIA INTELLIGENCE', 1)
        for mention in agent3_result['twitter_intelligence'][:3]:
            doc.add_paragraph(f'Date: {mention.get("date", "Unknown")}')
            doc.add_paragraph(f'Source: {mention.get("author", "Unknown")}')
            doc.add_paragraph(f'Summary: {mention.get("summary", "No summary available")}')
            doc.add_paragraph(f'Strategic Relevance: {mention.get("strategic_relevance", "Unknown")}')
            doc.add_paragraph('')

    # Executive Intelligence
    if agent3_result.get('executive_intelligence'):
        doc.add_heading('EXECUTIVE INTELLIGENCE', 1)
        for interview in agent3_result['executive_intelligence'][:3]:
            doc.add_paragraph(f'Date: {interview.get("date", "Unknown")}')
            doc.add_paragraph(f'Source: {interview.get("source", "Unknown")}')
            doc.add_paragraph(f'Executive: {interview.get("executive", "Unknown")}')
            doc.add_paragraph(f'Title: {interview.get("title", "No title")}')
            doc.add_paragraph(f'Key Insights: {interview.get("key_insights", "No insights available")}')
            doc.add_paragraph(f'Consulting Relevance: {interview.get("consulting_relevance", "Unknown")}')
            doc.add_paragraph('')

    # Intelligence Dashboard
    dashboard = result.get('intelligence_dashboard', {})
    if dashboard:
        doc.add_heading('INTELLIGENCE DASHBOARD', 1)

        overall_assessment = dashboard.get('overall_assessment', {})
        engagement_readiness = overall_assessment.get('engagement_readiness', 'Unknown')
        doc.add_paragraph(f'Engagement Readiness: {engagement_readiness}')

        next_steps = overall_assessment.get('recommended_next_steps', [])
        if next_steps:
            doc.add_paragraph('Recommended Next Steps:')
            for i, step in enumerate(next_steps, 1):
                doc.add_paragraph(f'{i}. {step}', style='List Number')

    # Footer
    doc.add_paragraph('')
    doc.add_paragraph('Generated by EPAM Business Intelligence System', style='Intense Quote')
    doc.add_paragraph(f'Report generated on {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}', style='Intense Quote')

    return doc

@app.route('/api/analyze', methods=['POST'])
def api_analyze():
    """API endpoint for programmatic analysis"""
    try:
        data = request.get_json()
        ticker = data.get('ticker', '').strip().upper()

        if not ticker:
            return jsonify({'error': 'Ticker symbol required'}), 400

        grok_key = data.get('grok_key')
        perplexity_key = data.get('perplexity_key')

        # Run analysis
        master_chain = MasterEPAMIntelligenceChain(grok_key, perplexity_key)
        result = master_chain.analyze_company_complete(ticker)

        return jsonify(result)

    except Exception as e:
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    # Check if templates directory exists
    templates_dir = os.path.join(os.path.dirname(__file__), 'templates')
    if not os.path.exists(templates_dir):
        os.makedirs(templates_dir)

    print("🚀 Starting EPAM Business Intelligence Web UI")
    print("📱 Access the application at: http://localhost:8080")
    print("🎯 Enter a ticker symbol to analyze company opportunities")
    print("📄 Generate and download comprehensive Word reports")

    app.run(debug=True, host='0.0.0.0', port=8080)